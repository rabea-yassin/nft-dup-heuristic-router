"""Build the project poster as an editable one-slide PowerPoint file.

    python3 report/figures/make_figures.py   # charts -> poster/assets/*.png
    python3 poster/make_poster.py            # -> poster/NFT_Copymint_Poster.pptx

The slide is 70 x 100 cm portrait. Every number on it comes from the final
report (report/), which cites PROGRESS.md for each one. Export to PDF from
PowerPoint (File > Export > PDF) to get poster/NFT_Copymint_Poster.pdf.
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
DOCS = HERE.parent / "docs"
OUT = HERE / "NFT_Copymint_Poster.pptx"

# Frame palettes: DARK fills the banner and footer, HEAD is the main text on it
# and LIGHT the secondary text, BAR fills the section bars, TITLE is the
# section-title text, ACCENT is the banner underline and text accents, STRIPE
# the bar's left stripe. HEAD defaults to white, STRIPE to ACCENT.
# Chosen with --palette.
PALETTES = {
    "navy":       dict(DARK="1B3A6B", LIGHT="C9D6EA", BAR="E3EAF5", ACCENT="2A78D6", TITLE="1B3A6B",
                       STRIPE="1B3A6B"),
    "beige":      dict(DARK="D8C7A6", HEAD="3A2E1F", LIGHT="5C4A33", BAR="EFE5D2", ACCENT="8B6B3E",
                       TITLE="3A2E1F"),
    "wine":       dict(DARK="7A2E39", LIGHT="F1D9DC", BAR="F7ECEC", ACCENT="A33A47", TITLE="4A1E25"),
    "espresso":   dict(DARK="4A3426", LIGHT="E6D8C8", BAR="F4EDE3", ACCENT="B5793E", TITLE="3A281C"),
    "terracotta": dict(DARK="8C3B24", LIGHT="F2D5C6", BAR="FBEDE5", ACCENT="C2643F", TITLE="5A2616"),
}
DEFAULT_PALETTE = "navy"   # Technion navy; the others were alternatives considered
DARK = HEAD = BAR = ACCENT = STRIPE = LIGHT = TITLE = None   # set by use_palette()


def use_palette(name):
    colours = dict(HEAD="FFFFFF")
    colours.update(PALETTES[name])
    colours.setdefault("STRIPE", colours["ACCENT"])
    for key, hexval in colours.items():
        globals()[key] = RGBColor.from_string(hexval)


ORB = RGBColor(0x2A, 0x78, 0xD6)
SHASH = RGBColor(0xEB, 0x68, 0x34)
INK = RGBColor(0x2B, 0x2B, 0x2B)
INK_2 = RGBColor(0x5F, 0x63, 0x68)
PALE = RGBColor(0xEE, 0xF3, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

W, H = 70.0, 100.0          # slide size, cm
MARGIN = 2.0
GUTTER = 1.6
COL_W = (W - 2 * MARGIN - GUTTER) / 2
COL_X = (MARGIN, MARGIN + COL_W + GUTTER)
BODY = 26                   # body text size, pt
SECTION_GAP = 1.3


# --------------------------------------------------------------------------- #
# Small drawing helpers (all positions in cm)                                 #
# --------------------------------------------------------------------------- #

def text(slide, x, y, w, h, paras, size=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bold=False, spacing=1.12, space_after=6):
    """Add a text box. `paras` is a list of paragraphs; each paragraph is a
    string or a list of (text, style) runs, style being a dict with optional
    bold / color / size / italic."""
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, style in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = FONT
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.italic = style.get("italic", False)
            f.color.rgb = style.get("color", color)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, width=1.5):
    s = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def label_box(slide, x, y, w, h, label, fill=WHITE, line=INK_2, size=22,
              color=INK, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, width=1.5):
    s = rect(slide, x, y, w, h, fill, line, shape, width)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    tf = s.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Cm(0.1))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = label.split("\n")
    for i, line_text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line_text
        r.font.name = FONT
        r.font.size = Pt(size if i == 0 else size - 5)
        r.font.bold = bold and i == 0
        r.font.color.rgb = color if i == 0 else INK_2
    return s


def arrow(slide, x1, y1, x2, y2, color=INK_2, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def line(slide, x1, y1, x2, y2, color=INK_2, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def picture(slide, path, x, y, w=None, h=None):
    """Place an image scaled to width w (or height h), keeping its aspect ratio.
    Returns the placed height in cm."""
    iw, ih = Image.open(path).size
    if w is not None:
        h = w * ih / iw
    else:
        w = h * iw / ih
    slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))
    return h


def header(slide, col, y, title):
    """A tinted section bar with an accent stripe; returns the y below it."""
    x = COL_X[col]
    bar = rect(slide, x, y, COL_W, 2.3, BAR)
    rect(slide, x, y, 0.45, 2.3, STRIPE)
    tf = bar.text_frame
    tf.margin_left = Cm(1.2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(38)
    r.font.bold = True
    r.font.color.rgb = TITLE
    return y + 2.3 + 0.8


def B(t, **kw):
    return (t, dict(bold=True, **kw))


def N(t, **kw):
    return (t, dict(kw))


# --------------------------------------------------------------------------- #
# The poster                                                                  #
# --------------------------------------------------------------------------- #

def banner(slide):
    rect(slide, 0, 0, W, 12.6, DARK)
    rect(slide, 0, 12.6, W, 0.3, ACCENT)
    text(slide, MARGIN, 1.0, 30, 2.4,
         [[B("TECHNION", size=30, color=HEAD)],
          [N("Israel Institute of Technology", size=20, color=LIGHT)]],
         space_after=0, spacing=1.0)
    text(slide, W - MARGIN - 30, 1.0, 30, 2.4,
         [[N("02360340  ·  Project in Computer Communications", size=20, color=LIGHT)]],
         align=PP_ALIGN.RIGHT)
    text(slide, MARGIN, 3.6, W - 2 * MARGIN, 3.6,
         [[B("Detecting Copied NFT Images", size=80, color=HEAD)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    text(slide, MARGIN, 7.3, W - 2 * MARGIN, 2.0,
         [[N("Replacing a broken signal in an NFT copymint detector",
             size=40, color=LIGHT)]],
         align=PP_ALIGN.CENTER, space_after=0)
    text(slide, MARGIN, 10.0, W - 2 * MARGIN, 1.6,
         [[N("Rabea Yassin  ·  Ahmad Nassar        Instructor: Eran Tavor",
             size=28, color=HEAD)]],
         align=PP_ALIGN.CENTER, space_after=0)


def left_column(slide):
    col, x = 0, COL_X[0]
    y = 15.0

    # 1. The problem -------------------------------------------------------- #
    y = header(slide, col, y, "The problem: copymints")
    text(slide, x, y, COL_W, 4.4, [
        [N("A copymint takes someone's NFT artwork, edits it a little, and mints it "
           "as new. Exact hashing (SHA-256) can't see it: one changed pixel gives "
           "an unrelated digest. Detection needs fingerprints of "),
         B("what the image looks like"), N(".")]])
    y += 4.7
    img_w, gap = 7.0, 1.4
    x0 = x + (COL_W - 3 * img_w - 2 * gap) / 2
    for i, (name, cap) in enumerate([("example_original.png", "original"),
                                      ("example_rotated.png", "rotated copy"),
                                      ("example_recoloured.png", "recoloured copy")]):
        picture(slide, ASSETS / name, x0 + i * (img_w + gap), y, w=img_w)
        text(slide, x0 + i * (img_w + gap), y + img_w + 0.2, img_w, 1.2, [cap],
             size=22, color=INK_2, align=PP_ALIGN.CENTER)
    y += img_w + 1.6 + SECTION_GAP

    # 2. The paper's detector ---------------------------------------------- #
    y = header(slide, col, y, "What the researchers built")
    text(slide, x, y, COL_W, 2.4, [
        [N("Kotzer et al. compute four perceptual hashes per image, search each in "
           "its own BK-tree, and call it a copy when "), B("at least two agree"),
         N(".")]])
    y += 3.0
    names = [("aHash", "brightness layout"), ("pHash", "coarse shapes"),
             ("hsvHash", "colour mix"), ("sHash", "image pieces (cropping)")]
    bw, bh, bgap = 10.5, 2.3, 0.45
    vx = x + bw + 6.0
    for i, (n, d) in enumerate(names):
        by = y + i * (bh + bgap)
        is_s = n == "sHash"
        label_box(slide, x + 1.0, by, bw, bh, f"{n}\n{d}",
                  line=SHASH if is_s else INK_2, width=3.5 if is_s else 1.5)
        arrow(slide, x + 1.0 + bw, by + bh / 2, vx, y + 2 * (bh + bgap) - bgap / 2)
    vy = y + 2 * (bh + bgap) - bgap / 2 - 2.0
    label_box(slide, vx, vy, 7.0, 4.0, "2 of 4\nagree?", fill=PALE, size=28, bold=True)
    arrow(slide, vx + 7.0, vy + 2.0, vx + 9.0, vy + 2.0)
    label_box(slide, vx + 9.0, vy + 0.75, 5.6, 2.5, "copy /\nnot a copy", size=22)
    y += 4 * (bh + bgap) + SECTION_GAP

    # 3. The sHash finding -------------------------------------------------- #
    y = header(slide, col, y, "What we found: sHash isn't a distance")
    text(slide, x, y, COL_W, 4.0, [
        [N("sHash fingerprints each "), B("piece"), N(" of an image and compares the "
           "lists "), B("one way"), N(". In the data the authors shared with us, the "
           "reported distance matches one direction on all 1,802 pairs, and the other "
           "on only 240.")]])
    y += 4.4
    # Triangle: A (cat) - B (cat+dog) - C (dog).
    r = 2.0
    ax_, bx_, cx_ = x + 1.5, x + 9.5, x + 17.5
    top, bot = y + 0.2, y + 4.4
    nodes = [(ax_, bot, "A\ncat"), (bx_, top, "B\ncat+dog"), (cx_, bot, "C\ndog")]
    line(slide, ax_ + r, bot + r, bx_ + r, top + r, width=3)
    line(slide, bx_ + r, top + r, cx_ + r, bot + r, width=3)
    line(slide, ax_ + r, bot + r, cx_ + r, bot + r, color=SHASH, width=4)
    for nx, ny, lab in nodes:
        label_box(slide, nx, ny, 2 * r, 2 * r, lab, shape=MSO_SHAPE.OVAL, size=22,
                  bold=True)
    text(slide, (ax_ + bx_) / 2 + 0.2, (top + bot) / 2 + 0.6, 3, 1.3, ["2"], size=30,
         bold=True)
    text(slide, (bx_ + cx_) / 2 + 2.6, (top + bot) / 2 + 0.6, 3, 1.3, ["16"], size=30,
         bold=True)
    text(slide, (ax_ + cx_) / 2 + 0.8, bot + 2 * r + 0.1, 4, 1.3, ["30"], size=30,
         bold=True, color=SHASH)
    text(slide, x + 23.0, y + 2.4, COL_W - 23.0, 3.0,
         [[B("30 > 2 + 16", size=40, color=SHASH)]], anchor=MSO_ANCHOR.MIDDLE)
    y += 10.2
    text(slide, x, y, COL_W, 3.0, [
        [N("A search tree assumes the direct route is never longer than the detour. "
           "With sHash it can be, so the tree "),
         B("skips a branch holding a real copy"), N(" and nothing looks wrong.")]])
    y += 3.4 + SECTION_GAP

    # 4. The cost ----------------------------------------------------------- #
    y = header(slide, col, y, "The honest part: what ORB costs")
    text(slide, x, y, COL_W, 4.0, [
        [N("The paper stores its four hashes, "), B("≈56 bytes"), N(", inside each "
           "mint transaction (300–500 bytes). ORB stores "), B("14,560 bytes"),
         N(" per image: about 36 transactions.")]])
    y += 4.2
    h = picture(slide, ASSETS / "fig_storage.png", x + 6.5, y, w=COL_W - 13.0)
    y += h + 0.4
    text(slide, x, y, COL_W, 2.6, [
        [N("Shrink it and the advantage goes first: at 32 landmarks (1 KB, still 2.6 "
           "transactions) ORB only ties sHash. "),
         B("No size both fits and wins.")]])
    y += 2.8
    print(f"left column ends at {y:.1f} cm (footer at {H - 2.4:.1f})")


def right_column(slide):
    col, x = 1, COL_X[1]
    y = 15.0

    def side_by_side(img, img_w, paras, size=24):
        """Image on the left, its explanation beside it; returns the block height."""
        h = picture(slide, img, x, y, w=img_w)
        text(slide, x + img_w + 0.8, y, COL_W - img_w - 0.8, h, paras, size=size,
             anchor=MSO_ANCHOR.MIDDLE)
        return h

    # 5. ORB ---------------------------------------------------------------- #
    y = header(slide, col, y, "Our replacement: ORB feature matching")
    y += side_by_side(DOCS / "orb_match_rotate.png", 17.5, [
        [B("Find landmarks"), N(" (corners) in both images, "), B("match"),
         N(" them, and let "), B("RANSAC"),
         N(" keep only the matches that agree on one transformation.")],
        [N("Here 204 matches agree on one rotation. ORB replaces sHash; the other "
           "three hashes and the vote are untouched.")]]) + SECTION_GAP

    # 6. Result 1 ----------------------------------------------------------- #
    y = header(slide, col, y, "Result 1: +14.2 F1 on sHash's own job")
    y += side_by_side(ASSETS / "fig_head_to_head.png", 18.5, [
        [N("On the crop and rotation copies sHash was built for.")],
        [N("sHash was "), B("allowed to pick its threshold on the test answers"),
         N("; ORB's was set on training data.")],
        [N("sHash keeps its recall by flagging "), B("1,630 of 2,400"),
         N(" innocent images.")]]) + SECTION_GAP

    # 7. Result 2 ----------------------------------------------------------- #
    y = header(slide, col, y, "Result 2: it wins almost everywhere")
    y += side_by_side(ASSETS / "fig_by_edit_type.png", 20.5, [
        [N("Share of each edit's copies ORB catches (precision 91.2%).")],
        [N("One clean failure: "), B("pixelation"),
         N(" destroys the detail corners are made of.")],
        [N("*Always includes rotation; pure mirrors: 100%.", size=20, color=INK_2)]],
        size=23) + SECTION_GAP

    # 8. Result 3 ----------------------------------------------------------- #
    y = header(slide, col, y, "Result 3: +8.9 F1 in the whole detector")
    y += side_by_side(ASSETS / "fig_whole_detector.png", 18.5, [
        [N("Change only sHash; keep the paper's own thresholds.")],
        [N("At the same "), B("98.1% precision"), N(", recall rises from "),
         B("64.1% to 77.2%"), N(": ~1,500 more copies caught.")],
        [N("Our test set is 82.6% copies, so equal precision is the fair "
           "comparison.", size=20, color=INK_2)]]) + SECTION_GAP

    # 9. Lessons ------------------------------------------------------------ #
    y = header(slide, col, y, "What we learned")
    text(slide, x, y, COL_W, 11.0, [
        [B("Check the maths before optimising it. "),
         N("We set out to make the search faster and found it was unsound.")],
        [B("How a signal is stored decides how it can be searched. "),
         N("A fingerprint fits a tree; a set of landmarks needs LSH, and far more "
           "space.")],
        [B("A negative result is still a result. "),
         N("Our edit classifier predicts which signals an edit breaks (100% / 97.5%), "
           "but doesn't improve the detector: broken signals go silent rather than "
           "voting wrongly.")],
        [B("Next: ", color=ACCENT),
         N("keep ORB's landmarks off-chain with a small commitment on-chain. Letting the "
           "cheap hashes shortlist candidates for an off-chain ORB check is an untested "
           "idea.")],
    ], size=24, space_after=10)
    y += 13.0
    print(f"right column ends at {y:.1f} cm (footer at {H - 2.4:.1f})")


def footer(slide):
    rect(slide, 0, H - 2.4, W, 2.4, DARK)
    text(slide, MARGIN, H - 2.0, W - 2 * MARGIN, 1.6, [
        [N("Based on Kotzer, Reviriego, Conde Diaz and Rottenstreich, \"Combating NFT "
           "Copymints in Blockchain Networks: An Image Hashing Approach\".   "
           "Code and full report in the project repository.", size=20, color=LIGHT)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main():
    import argparse
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--palette", choices=sorted(PALETTES), default=DEFAULT_PALETTE)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    use_palette(args.palette)

    prs = Presentation()
    prs.slide_width = Cm(W)
    prs.slide_height = Cm(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    banner(slide)
    left_column(slide)
    right_column(slide)
    footer(slide)
    prs.save(args.out)
    print("wrote", args.out)


if __name__ == "__main__":
    main()
